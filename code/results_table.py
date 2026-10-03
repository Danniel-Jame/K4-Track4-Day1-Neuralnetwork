"""results_table.py — Bản hoàn thiện.

Nhiệm vụ: lưu kết quả từng lần chạy ra JSON, rồi điền vào experiments.xlsx từ mẫu
templates/experiment_table_template.xlsx.

Tên cột của sheet "Experiments" (giữ nguyên, đúng thứ tự mẫu):
    exp_id, group, description, loss, optimizer, lr, weight_decay, batch, epochs, hidden, dropout,
    clip_norm, precision, init, seed, step0_loss, best_val_loss, best_epoch, final_train_loss,
    final_val_loss, val_acc, val_macro_f1, time_per_epoch_s, peak_mem_MB, diverged,
    eval_acc, eval_macro_f1, figure_file, notes
"""
from __future__ import annotations

import json
import os
from pathlib import Path
import openpyxl


def save_result(result: dict, results_dir: str = "../results") -> str:
    """Ghi result["cfg"], result["history"], result["summary"] (KHÔNG ghi best_state) ra
    <results_dir>/<exp_id>.json. Trả về đường dẫn file. Tạo thư mục nếu chưa có.
    """
    exp_id = result.get("exp_id") or result.get("cfg", {}).get("exp_id", "unknown")
    os.makedirs(results_dir, exist_ok=True)
    out_path = os.path.join(results_dir, f"{exp_id}.json")

    # Loại bỏ best_state (mô hình weights) để giữ file JSON nhẹ và đúng định dạng
    clean_result = {
        "exp_id": exp_id,
        "cfg": result.get("cfg", {}),
        "history": result.get("history", {}),
        "summary": result.get("summary", {})
    }

    with open(out_path, "w", encoding="utf-8") as f:
        json.dump(clean_result, f, ensure_ascii=False, indent=2)

    return out_path


def load_results(results_dir: str = "../results") -> list[dict]:
    """Đọc mọi file *.json trong results_dir, trả về danh sách dict (sắp theo exp_id)."""
    results_path = Path(results_dir)
    if not results_path.exists():
        return []

    results = []
    for p in sorted(results_path.glob("*.json")):
        try:
            with open(p, "r", encoding="utf-8") as f:
                res = json.load(f)
                results.append(res)
        except Exception as e:
            print(f"Cảnh báo: Không thể đọc file JSON {p}: {e}")

    # Sắp xếp danh sách theo exp_id
    results.sort(key=lambda x: x.get("exp_id", x.get("cfg", {}).get("exp_id", "")))
    return results


def to_row(result: dict, eval_scores: dict | None = None, notes: str = "") -> dict:
    """Biến một kết quả thành một dòng của bảng: gộp cfg + summary (+ eval_acc, eval_macro_f1 nếu có)
    + figure_file = f"figures/{exp_id}.png". Khoá phải trùng tên cột ở đầu file.
    Chỉ truyền eval_scores cho baseline và cấu hình cuối cùng.
    """
    cfg = result.get("cfg", {})
    summary = result.get("summary", {})
    exp_id = result.get("exp_id") or cfg.get("exp_id", "unknown")

    # Xử lý dạng biểu diễn của cấu trúc hidden (ví dụ: (256, 128) -> "256-128")
    hidden_val = cfg.get("hidden", (256, 128))
    if isinstance(hidden_val, (list, tuple)):
        hidden_str = "-".join(str(h) for h in hidden_val)
    else:
        hidden_str = str(hidden_val)

    row = {
        "exp_id": exp_id,
        "group": cfg.get("group", "baseline"),
        "description": cfg.get("description", ""),
        "loss": cfg.get("loss", "ce"),
        "optimizer": cfg.get("optimizer", "sgd_momentum"),
        "lr": cfg.get("lr", 0.01),
        "weight_decay": cfg.get("weight_decay", 0.0),
        "batch": cfg.get("batch_size", 512),
        "epochs": cfg.get("epochs", 20),
        "hidden": hidden_str,
        "dropout": cfg.get("dropout", 0.0),
        "clip_norm": cfg.get("clip_norm", None) or "none",
        "precision": cfg.get("precision", "fp32"),
        "init": cfg.get("init", "he"),
        "seed": cfg.get("seed", 42),
        "step0_loss": summary.get("step0_loss", None),
        "best_val_loss": summary.get("best_val_loss", None),
        "best_epoch": summary.get("best_epoch", None),
        "final_train_loss": summary.get("final_train_loss", None),
        "final_val_loss": summary.get("final_val_loss", None),
        "val_acc": summary.get("val_acc", None),
        "val_macro_f1": summary.get("val_macro_f1", None),
        "time_per_epoch_s": summary.get("time_per_epoch_s", None),
        "peak_mem_MB": summary.get("peak_mem_MB", None),
        "diverged": summary.get("diverged", False),
        "eval_acc": eval_scores.get("accuracy", None) if eval_scores else None,
        "eval_macro_f1": eval_scores.get("macro_f1", None) if eval_scores else None,
        "figure_file": f"figures/{exp_id}.png",
        "notes": notes,
    }

    return row


def write_xlsx(rows: list[dict], template_path: str, out_path: str) -> None:
    """Điền các dòng vào sheet "Experiments" của mẫu, từ dòng 2 trở xuống, rồi lưu thành out_path.

    Các bước (openpyxl):
      1. wb = openpyxl.load_workbook(template_path)   # KHÔNG dùng data_only=True (sẽ mất công thức)
      2. ws = wb["Experiments"]; đọc tiêu đề dòng 1 để biết cột nào ứng với khoá nào
      3. với mỗi row: ghi giá trị vào đúng cột; BỎ QUA các cột công thức
      4. wb.save(out_path)
    """
    if not os.path.exists(template_path):
        raise FileNotFoundError(f"Không tìm thấy file mẫu tại '{template_path}'")

    # Load workbook mà KHÔNG bật data_only=True để giữ nguyên các công thức Excel
    wb = openpyxl.load_workbook(template_path)
    if "Experiments" not in wb.sheetnames:
        raise ValueError("Sheet 'Experiments' không tồn tại trong template Excel.")

    ws = wb["Experiments"]

    # Đọc tiêu đề dòng 1 để xác định vị trí index của các cột
    col_map = {}
    for col_idx in range(1, ws.max_column + 1):
        header_val = ws.cell(row=1, column=col_idx).value
        if header_val:
            col_map[str(header_val).strip()] = col_idx

    # Danh sách cột chứa công thức cần BỎ QUA (không ghi đè)
    formula_cols = {
        "step0_gap_vs_lnC",
        "gap_val_minus_train",
        "delta_val_f1_vs_base",
        "beyond_noise",
    }

    # Điền các dòng từ dòng 2 trở xuống
    for row_offset, row_data in enumerate(rows):
        current_row = 2 + row_offset
        for key, val in row_data.items():
            if key in col_map and key not in formula_cols:
                col_idx = col_map[key]
                ws.cell(row=current_row, column=col_idx, value=val)

    # Đảm bảo thư mục đầu ra tồn tại trước khi lưu
    os.makedirs(os.path.dirname(os.path.abspath(out_path)), exist_ok=True)
    wb.save(out_path)
    print(f"Đã cập nhật file Excel thành công tại: {out_path}")